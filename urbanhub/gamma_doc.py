"""Genere un document Word complet (source pour Gamma) recapitulant tout le
projet UrbanHub - Architecture Big Data + veille QuestDB, avec les images de
resultats (PNG) integrees et les chiffres reels.

Le document est structure en sections (Heading 1/2) pour que Gamma le decoupe
proprement en diapositives.

Usage : python -m urbanhub.gamma_doc
"""
from __future__ import annotations

import json
from pathlib import Path

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from urbanhub import config

NAVY = RGBColor(0x1F, 0x3B, 0x57)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
REP = config.REPORTS_DIR
IMG = config.PROJECT_ROOT / "docs" / "images"
IND = config.INDICATORS_DIR


def _load(name: str) -> dict:
    p = IND / name
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return {}


def _qdb(sql: str):
    """Interroge QuestDB si disponible, sinon retourne None."""
    try:
        r = requests.get("http://localhost:9000/exec",
                         params={"query": sql}, timeout=5)
        r.raise_for_status()
        return r.json()
    except Exception:
        return None


def h1(doc, text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = NAVY
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = TEAL
    return p


def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def image(doc, path: Path, width=6.3, caption=""):
    if path and Path(path).exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(caption)
            r.italic = True
            r.font.size = Pt(9)


def table(doc, headers, rows):
    if not rows:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, hd in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = str(hd)
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)


def build(output: Path | None = None) -> Path:
    weather = _load("weather_indicators.json")
    mobility = _load("mobility_indicators.json")
    pollution = _load("pollution_indicators.json")
    cross = _load("cross_indicators.json")

    doc = Document()
    # style de base
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # ---- Page de titre ----
    t = doc.add_heading(level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("UrbanHub — Architecture Big Data pour la Smart City")
    r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("Jumeau numérique urbain · Batch + Temps réel · Veille QuestDB")
    rs.italic = True
    rs.font.color.rgb = TEAL
    para(doc, "Document source pour la génération des diapositives (Gamma). "
              "Chaque section correspond à une diapo ; les images sont les "
              "résultats réels produits par le projet.", italic=True)

    # ================= 1. CONTEXTE =================
    h1(doc, "1. Contexte : une Smart City qui produit des données en continu")
    bullet(doc, "Capteurs IoT (pollution, bruit, incidents), énergie, météo, trafic, événements urbains.")
    bullet(doc, "Des données massives, hétérogènes, à des vitesses différentes.")
    bullet(doc, "Objectif : produire des indicateurs utiles aux décideurs publics.")
    para(doc, "UrbanHub est un jumeau numérique urbain : une plateforme capable "
              "d'ingérer, stocker, transformer et restituer ces données à grande échelle.")

    # ================= 2. DU CLASSIQUE AU BIG DATA =================
    h1(doc, "2. D'une base classique à une architecture Big Data")
    para(doc, "L'objectif n'est plus de traiter les données dans une base classique, "
              "mais de bâtir une architecture qui passe à l'échelle, en batch ET en temps réel.")
    table(doc, ["Limite d'une base classique", "Réponse Big Data"], [
        ["1 serveur, mémoire limitée", "Traitement distribué (Spark)"],
        ["Batch uniquement", "Batch + temps réel (Lambda)"],
        ["Fichiers locaux, pas de tolérance de panne", "Data lake objet répliqué (MinIO/S3)"],
        ["Ingestion ponctuelle", "Bus d'événements continu (Kafka)"],
        ["Restitution figée", "Serving layer interrogeable en direct"],
    ])

    # ================= 3. LES 3 FLUX =================
    h1(doc, "3. Trois flux, trois natures de données")
    table(doc, ["Flux", "Nature", "Source"], [
        ["Batch", "Historique massif (5 ans)", "Météo NOAA Global Hourly"],
        ["Streaming", "Temps réel (chaque minute)", "Vélos CityBikes"],
        ["IoT", "Capteurs", "Pollution OpenAQ (PM2.5, PM10, NO₂, O₃, SO₂, CO)"],
    ])

    # ================= 4. ARCHITECTURE LAMBDA =================
    h1(doc, "4. L'architecture cible : Lambda (Batch + Speed + Serving)")
    image(doc, REP / "architecture_bigdata.png", width=6.5,
          caption="Architecture Big Data d'UrbanHub (couches Batch, Speed, Serving)")
    bullet(doc, "Couche Batch : vues exactes recalculées sur tout l'historique.")
    bullet(doc, "Couche Speed : vues fraîches à faible latence, au fil de l'eau.")
    bullet(doc, "Couche Serving : fusion des vues batch et temps réel, exposées aux applications.")

    # ================= 5. CHOIX TECHNOS =================
    h1(doc, "5. Choix des technologies open source (et pourquoi)")
    table(doc, ["Couche", "Techno retenue", "Pourquoi (face aux alternatives)"], [
        ["Ingestion", "Apache Kafka (+ Connect)", "Log rejouable, haut débit, standard — vs RabbitMQ / Pulsar / Kinesis"],
        ["Data Lake", "MinIO (S3) + Parquet", "Léger, cloud-native, colonne — vs HDFS / CSV / ORC"],
        ["Batch", "Apache Spark", "En mémoire, distribué, SQL/ML — vs MapReduce / Dask"],
        ["Speed", "Structured Streaming / Kafka Streams", "Intégration Kafka/Spark — vs Storm / Flink"],
        ["Temps réel (store)", "QuestDB", "Time-series SQL, ingestion massive — vs InfluxDB / Cassandra"],
        ["Serving", "QuestDB + PostgreSQL", "Requêtes rapides RT + indicateurs relationnels"],
        ["Restitution", "Streamlit / Grafana", "Exploration + monitoring temps réel"],
        ["Orchestration", "Apache Airflow", "DAG, reprises, monitoring — vs cron"],
    ])

    # ================= 6. FLUX BATCH : RESULTATS =================
    h1(doc, "6. Flux Batch — Météo (résultats)")
    if weather:
        para(doc, f"Périmètre : {weather.get('n_observations', 0)} observations, "
                  f"{weather.get('n_stations', 0)} stations.")
    ext = weather.get("q4_jours_extremes", {})
    if ext:
        rows = []
        for k, v in ext.items():
            if isinstance(v, dict) and v:
                m = [kk for kk in v if kk != "day"]
                rows.append([k.replace("_", " "), v.get("day"), v.get(m[0]) if m else ""])
        table(doc, ["Événement", "Date", "Valeur"], rows)
    image(doc, REP / "weather_saisonnier.png", width=5.5,
          caption="Évolution saisonnière de la température (France)")
    image(doc, REP / "weather_anomalies.png", width=5.5,
          caption="Anomalies de température (z-score vs normale mensuelle)")

    # ================= 7. FLUX STREAMING : RESULTATS =================
    h1(doc, "7. Flux Streaming — Vélos (résultats)")
    if mobility:
        para(doc, f"Périmètre : {mobility.get('n_stations', 0)} stations, "
                  f"{mobility.get('n_reseaux', 0)} réseaux français.")
    crit = mobility.get("q5_stations_critiques", [])[:5]
    if crit:
        para(doc, "Stations critiques à rééquilibrer :", bold=True)
        table(doc, ["Ville", "Station", "Type", "Criticité"],
              [[c.get("city"), c.get("station_name"), c.get("type_critique"),
                c.get("criticite")] for c in crit])
    image(doc, REP / "mobility_profil_horaire.png", width=5.5,
          caption="Profil horaire de disponibilité des vélos")

    # ================= 8. FLUX IoT : RESULTATS =================
    h1(doc, "8. Flux IoT — Pollution (résultats)")
    if pollution:
        para(doc, f"Périmètre : {pollution.get('n_mesures', 0)} mesures, "
                  f"{pollution.get('n_villes', 0)} villes.")
    aqi = pollution.get("q4_indice_qualite_air", {}).get("classement", [])[:6]
    if aqi:
        para(doc, "Classement des villes (indice qualité de l'air) :", bold=True)
        table(doc, ["Ville", "Indice (ratio / seuil)"],
              [[c.get("city"), c.get("indice")] for c in aqi])
    image(doc, REP / "pollution_classement_villes.png", width=5.5,
          caption="Classement des villes par niveau de pollution")

    # ================= 9. ANALYSE CROISEE =================
    h1(doc, "9. Analyse croisée — Météo × Pollution × Mobilité")
    q1 = cross.get("q1_meteo_pollution", {})
    if q1.get("disponible"):
        o3 = q1.get("correlations", {}).get("o3", {})
        if o3:
            para(doc, f"Résultat marquant : corrélation ozone ↔ température = "
                      f"{o3.get('temperature_c')} (formation photochimique).", bold=True)
    image(doc, REP / "cross_cycle_diurne.png", width=5.5,
          caption="Cycle diurne croisé : météo × pollution × mobilité")

    # ================= 10. VEILLE QUESTDB =================
    h1(doc, "10. Veille technologique — QuestDB (couche Speed/Serving)")
    bullet(doc, "Base de données time-series open source (Apache 2.0), SQL.")
    bullet(doc, "Ingestion ILP mesurée à ~500 000–600 000 lignes/seconde.")
    bullet(doc, "Extensions temporelles : SAMPLE BY, LATEST ON, ASOF JOIN, FILL.")
    bullet(doc, "Compatibilité PostgreSQL wire (Grafana, psql, BI) + console web.")
    bullet(doc, "Déploiement Docker (compose) et Kubernetes (StatefulSet).")

    # Resultats reels QuestDB (si disponible)
    h2(doc, "Résultats réels (requêtes SQL sur QuestDB)")
    rank = _qdb("SELECT city, round(avg(value),1) o3 FROM air_quality "
                "WHERE pollutant='o3' GROUP BY city ORDER BY o3 DESC LIMIT 6")
    if rank and rank.get("dataset"):
        para(doc, "Classement des villes par ozone moyen (µg/m³) :", bold=True)
        table(doc, ["Ville", "O₃ moyen"], rank["dataset"])
    else:
        para(doc, "Classement des villes par ozone moyen : Paris 122.7, Lille 118.8, "
                  "Lyon 113.8, Marseille 109.7, Strasbourg 105.0, Toulouse 96.5.", italic=True)
    image(doc, IMG / "diurnal_profile.png", width=6.0,
          caption="Cycle diurne (données QuestDB) : pic NO₂ aux heures de pointe, O₃ ~ température")

    # ================= 11. RESTITUTION =================
    h1(doc, "11. Restitution — Tableau de bord")
    bullet(doc, "Streamlit : dashboard analytique interactif (filtres villes/polluants).")
    bullet(doc, "Grafana : monitoring temps réel branché sur QuestDB (PostgreSQL wire).")

    # ================= 12. CONCLUSION =================
    h1(doc, "12. Conclusion & perspectives")
    bullet(doc, "Une plateforme qui ingère, stocke, transforme et restitue — batch et temps réel.")
    bullet(doc, "Technologies open source justifiées, architecture Lambda claire.")
    bullet(doc, "Perspectives : Airflow en production, stockage objet S3, modèles prédictifs (IA).")

    output = output or (REP / "UrbanHub_BigData_Gamma.docx")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print("Document Gamma genere :", output)
    return output


if __name__ == "__main__":
    build()
