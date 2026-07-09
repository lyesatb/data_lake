"""Generateur du document Word (.docx) des reponses aux questions metier.

Compile, pour les 4 parties du projet, chaque question metier avec la methode
employee et la reponse chiffree, a partir des indicateurs reels du data lake.

Usage :
    python -m urbanhub.cli doc
    # ou
    python -m urbanhub.report_doc
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from urbanhub import config
from urbanhub.utils.logging_utils import get_logger

log = get_logger("report_doc")

NAVY = RGBColor(0x1F, 0x3B, 0x57)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)


def _load(name: str) -> dict:
    p = config.INDICATORS_DIR / name
    if p.exists():
        with p.open(encoding="utf-8") as fh:
            return json.load(fh)
    return {}


def _heading(doc, text, level=1, color=NAVY):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.color.rgb = color
    return h


def _question(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("❓ " + text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = TEAL
    return p


def _kv(doc, label, value):
    p = doc.add_paragraph()
    r = p.add_run(f"{label} : ")
    r.bold = True
    p.add_run(str(value))
    return p


def _para(doc, text, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    return p


def _table(doc, headers, rows):
    if not rows:
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = str(h)
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    doc.add_paragraph()


def build(output: Path | None = None) -> Path:
    weather = _load("weather_indicators.json")
    mobility = _load("mobility_indicators.json")
    pollution = _load("pollution_indicators.json")
    cross = _load("cross_indicators.json")

    doc = Document()

    # -- Titre ------------------------------------------------------------- #
    title = doc.add_heading(level=0)
    r = title.add_run("UrbanHub — Réponses aux questions métier")
    r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Jumeau numérique urbain · Batch météo · Streaming vélos · IoT pollution")
    sr.italic = True
    sr.font.color.rgb = TEAL
    doc.add_paragraph()
    _para(doc, "Ce document reprend, pour chaque flux, les questions métier du "
               "sujet et y répond à partir des données réellement collectées, "
               "nettoyées et analysées par la plateforme. Les chiffres sont "
               "issus des indicateurs du data lake (data/curated).", italic=True)

    # ==================================================================== #
    # PARTIE 1 — METEO (BATCH)
    # ==================================================================== #
    _heading(doc, "Partie 1 — Flux Batch : analyse météorologique urbaine", 1)
    if weather:
        _para(doc, f"Périmètre : {weather.get('n_observations', 0):,} observations "
                   f"issues de {weather.get('n_stations', 0)} stations françaises "
                   f"(source NOAA Global Hourly).".replace(",", " "))

    _question(doc, "Peut-on identifier des périodes météorologiques anormales ?")
    q1 = weather.get("q1_periodes_anormales", {})
    _kv(doc, "Méthode", q1.get("definition", "z-score de la température vs normale mensuelle"))
    _kv(doc, "Réponse", f"Oui — {q1.get('nombre_jours_anormaux', 0)} jour(s) anormal(aux) détecté(s).")
    _table(doc, ["Date", "Température (°C)", "z-score"],
           [[a.get("day"), round(a.get("temperature_c", 0), 1), round(a.get("temp_zscore", 0), 2)]
            for a in q1.get("top_anomalies", [])[:5]])

    _question(doc, "Existe-t-il une corrélation entre conditions météo et visibilité ?")
    corr = weather.get("q2_correlation_visibilite", {})
    _kv(doc, "Méthode", "Coefficient de corrélation de Pearson entre la visibilité et les autres variables.")
    if corr:
        _table(doc, ["Variable", "Corrélation avec la visibilité"],
               [[k, round(v, 3)] for k, v in corr.items()])
        _para(doc, "Interprétation : la visibilité diminue quand l'humidité augmente "
                   "(brouillard) et augmente avec la température.", italic=True)

    _question(doc, "Quelle est l'évolution saisonnière de la température en France ?")
    season = weather.get("q3_temperature_saisonniere_france", {})
    if season:
        _table(doc, ["Saison", "Température moyenne (°C)"],
               [[k, v] for k, v in season.items()])
    villes = weather.get("q3_temperature_saisonniere_villes", [])
    if villes:
        _para(doc, "Détail par ville :")
        keys = [k for k in villes[0].keys() if k != "city"]
        _table(doc, ["Ville"] + keys,
               [[v.get("city")] + [v.get(k) for k in keys] for v in villes])

    _question(doc, "Quels sont les jours présentant des conditions météorologiques extrêmes ?")
    ext = weather.get("q4_jours_extremes", {})
    if ext:
        rows = []
        for k, v in ext.items():
            if isinstance(v, dict) and v:
                metric = [kk for kk in v if kk != "day"]
                rows.append([k.replace("_", " "), v.get("day"),
                             v.get(metric[0]) if metric else ""])
        _table(doc, ["Événement", "Date", "Valeur"], rows)

    # ==================================================================== #
    # PARTIE 2 — MOBILITE (STREAMING)
    # ==================================================================== #
    _heading(doc, "Partie 2 — Flux Streaming : mobilité urbaine (vélos)", 1)
    if mobility:
        _para(doc, f"Périmètre : {mobility.get('n_stations', 0):,} stations sur "
                   f"{mobility.get('n_reseaux', 0)} réseaux français "
                   f"(source CityBikes).".replace(",", " "))

    _question(doc, "Quelles stations présentent le plus fort taux d'utilisation ?")
    _kv(doc, "Méthode", "Taux d'occupation moyen (vélos disponibles / capacité) par station.")
    _table(doc, ["Ville", "Station", "Taux d'occupation", "Capacité"],
           [[s.get("city"), s.get("station_name"), s.get("occupancy_rate"), s.get("capacity")]
            for s in mobility.get("q1_stations_plus_utilisees", [])[:8]])

    _question(doc, "Peut-on identifier des zones où l'offre de vélos est insuffisante ?")
    _kv(doc, "Méthode", "Stations le plus souvent vides (part du temps sans vélo).")
    _table(doc, ["Ville", "Station", "% du temps vide", "Vélos dispo (moy.)"],
           [[s.get("city"), s.get("station_name"), s.get("pct_empty"), s.get("bikes_available")]
            for s in mobility.get("q2_offre_insuffisante", [])[:8]])

    _question(doc, "Quels sont les pics d'utilisation journaliers du réseau ?")
    _kv(doc, "Méthode", "Profil horaire de la disponibilité moyenne des vélos (0–23 h).")
    _para(doc, "Voir le graphique 'mobility_profil_horaire.png'. Le profil se "
               "précise à mesure que le flux temps réel tourne sur la journée.", italic=True)

    _question(doc, "Existe-t-il des déséquilibres géographiques dans la disponibilité ?")
    _kv(doc, "Méthode", "Agrégation par ville : occupation moyenne et part de stations vides/pleines.")
    _table(doc, ["Ville", "Stations", "Occupation moy.", "% vides", "% pleines"],
           [[d.get("city"), d.get("stations"), d.get("occupancy_moy"),
             d.get("pct_empty_moy"), d.get("pct_full_moy")]
            for d in mobility.get("q4_desequilibres_villes", [])[:8]])

    _question(doc, "Peut-on détecter des stations critiques nécessitant un rééquilibrage ?")
    _kv(doc, "Méthode", "Stations très souvent vides (pénurie) OU très souvent pleines (saturation), criticité ≥ 0.5.")
    _table(doc, ["Ville", "Station", "Type", "Criticité"],
           [[s.get("city"), s.get("station_name"), s.get("type_critique"), s.get("criticite")]
            for s in mobility.get("q5_stations_critiques", [])[:8]])

    # ==================================================================== #
    # PARTIE 3 — POLLUTION (IoT)
    # ==================================================================== #
    _heading(doc, "Partie 3 — Flux IoT : pollution urbaine", 1)
    if pollution:
        _para(doc, f"Périmètre : {pollution.get('n_mesures', 0):,} mesures sur "
                   f"{pollution.get('n_villes', 0)} villes · polluants : "
                   f"{', '.join(pollution.get('polluants', []))}.".replace(",", " ", 1))

    _question(doc, "Quelles villes sont les plus polluées (par polluant) ?")
    _kv(doc, "Méthode", "Concentration moyenne par ville et par polluant.")
    top = pollution.get("q1_ville_plus_polluee_par_polluant", {})
    if top:
        _table(doc, ["Polluant", "Ville la plus polluée", "Valeur (µg/m³)"],
               [[k, v.get("ville"), round(v.get("valeur", 0), 1)] for k, v in top.items()])

    _question(doc, "Quel est le profil horaire des polluants ?")
    _kv(doc, "Méthode", "Moyenne par heure de la journée (met en évidence les pics de trafic et d'ozone).")
    _para(doc, "Voir le graphique du profil horaire dans le tableau de bord (onglet Pollution).", italic=True)

    _question(doc, "Quels épisodes dépassent les seuils de qualité de l'air ?")
    dep = pollution.get("q3_depassements_seuils", {})
    _kv(doc, "Méthode", f"Comparaison aux seuils indicatifs OMS : {dep.get('seuils_ug_m3', {})}")
    _kv(doc, "Réponse", f"{dep.get('total_depassements', 0)} dépassements au total.")
    _table(doc, ["Ville", "Polluant", "Nb dépassements"],
           [[d.get("city"), d.get("parameter"), d.get("nb_depassements")]
            for d in dep.get("top", [])[:8]])

    _question(doc, "Comment classer les villes selon un indice de qualité de l'air ?")
    aqi = pollution.get("q4_indice_qualite_air", {})
    _kv(doc, "Méthode", aqi.get("definition", "Moyenne des ratios valeur/seuil (1.0 = niveau du seuil)."))
    _table(doc, ["Ville", "Indice (ratio moyen / seuil)"],
           [[c.get("city"), c.get("indice")] for c in aqi.get("classement", [])[:8]])

    # ==================================================================== #
    # PARTIE 4 — ANALYSE CROISEE
    # ==================================================================== #
    _heading(doc, "Partie 4 — Analyse croisée des données", 1)
    _para(doc, cross.get("methode", ""), italic=True)

    _question(doc, "Existe-t-il une relation entre conditions météo et pollution ?")
    q1c = cross.get("q1_meteo_pollution", {})
    if q1c.get("disponible"):
        _kv(doc, "Réponse", "Oui — corrélations calculées sur le cycle diurne.")
        corr = q1c.get("correlations", {})
        headers = ["Polluant"] + sorted({k for v in corr.values() for k in v})
        rows = []
        for pol, cc in corr.items():
            rows.append([pol] + [cc.get(h, "") for h in headers[1:]])
        _table(doc, headers, rows)
        o3 = corr.get("o3", {})
        if o3:
            _para(doc, f"Résultat marquant : l'ozone (O₃) est fortement corrélé à la "
                       f"température ({o3.get('temperature_c')}), signature de sa "
                       f"formation photochimique.", italic=True)
    else:
        _kv(doc, "Réponse", "Recouvrement temporel insuffisant sur l'échantillon de démonstration.")

    _question(doc, "Les conditions météo influencent-elles l'utilisation des vélos ?")
    q2c = cross.get("q2_meteo_velos", {})
    if q2c.get("disponible"):
        corr = q2c.get("correlations", {})
        rows = []
        for tgt, cc in corr.items():
            for k, v in cc.items():
                rows.append([tgt, k, v])
        _table(doc, ["Cible", "Variable météo", "Corrélation"], rows)
    else:
        _kv(doc, "Réponse", q2c.get("message", "Nécessite que le flux vélos ait tourné sur toute la journée."))

    _question(doc, "Peut-on identifier des conditions météo favorables à la mobilité douce ?")
    q3c = cross.get("q3_conditions_favorables", {})
    if q3c:
        _kv(doc, "Définition", q3c.get("definition_conditions_favorables", {}))
        _kv(doc, "Part d'heures favorables", f"{q3c.get('part_heures_favorables_pct', 0)} %")
        saison = q3c.get("part_favorable_par_saison_pct", {})
        if saison:
            _table(doc, ["Saison", "Part d'heures favorables (%)"],
                   [[k, v] for k, v in saison.items()])

    _question(doc, "Existe-t-il des périodes où mobilité, pollution et météo interagissent fortement ?")
    q4c = cross.get("q4_triple_interaction", {})
    _kv(doc, "Méthode", "Jointure des trois flux sur le cycle diurne (heure de la journée).")
    _kv(doc, "Heures communes analysées", q4c.get("n_heures_communes", q4c.get("n_points_communs", 0)))
    _para(doc, "Voir le graphique 'cross_cycle_diurne.png' (superposition température, "
               "NO₂ et vélos disponibles sur la journée).", italic=True)

    output = output or (config.REPORTS_DIR / "UrbanHub_reponses_questions_metier.docx")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    log.info("Document Q/R généré : %s", output)

    build_markdown()
    return output


# --------------------------------------------------------------------------- #
# Version Markdown (lisible directement sur GitHub)
# --------------------------------------------------------------------------- #
def _md_table(headers, rows):
    if not rows:
        return ""
    out = ["| " + " | ".join(str(h) for h in headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    for r in rows:
        out.append("| " + " | ".join("" if c is None else str(c) for c in r) + " |")
    return "\n".join(out) + "\n"


def build_markdown(output: Path | None = None) -> Path:
    weather = _load("weather_indicators.json")
    mobility = _load("mobility_indicators.json")
    pollution = _load("pollution_indicators.json")
    cross = _load("cross_indicators.json")
    L: list[str] = []

    L.append("# UrbanHub — Réponses aux questions métier\n")
    L.append("_Jumeau numérique urbain · Batch météo · Streaming vélos · IoT pollution._\n")
    L.append("Réponses issues des données réellement collectées et analysées "
             "(indicateurs du data lake `data/curated`).\n")

    # Partie 1
    L.append("## Partie 1 — Flux Batch : météo urbaine (NOAA)\n")
    _nobs = f"{weather.get('n_observations', 0):,}".replace(",", " ")
    L.append(f"Périmètre : {_nobs} observations, {weather.get('n_stations', 0)} stations.\n")
    q1 = weather.get("q1_periodes_anormales", {})
    L.append("### Peut-on identifier des périodes météo anormales ?\n")
    L.append(f"Méthode : {q1.get('definition', '')}.  ")
    L.append(f"Réponse : **oui**, {q1.get('nombre_jours_anormaux', 0)} jour(s) anormal(aux).\n")
    L.append(_md_table(["Date", "Température (°C)", "z-score"],
             [[a.get("day"), round(a.get("temperature_c", 0), 1), round(a.get("temp_zscore", 0), 2)]
              for a in q1.get("top_anomalies", [])[:5]]))
    L.append("### Corrélation météo ↔ visibilité ?\n")
    corr = weather.get("q2_correlation_visibilite", {})
    L.append(_md_table(["Variable", "Corrélation"], [[k, round(v, 3)] for k, v in corr.items()]))
    L.append("_La visibilité baisse quand l'humidité augmente (brouillard)._\n")
    L.append("### Évolution saisonnière de la température ?\n")
    L.append(_md_table(["Saison", "Température moyenne (°C)"],
             [[k, v] for k, v in weather.get("q3_temperature_saisonniere_france", {}).items()]))
    L.append("### Jours aux conditions extrêmes ?\n")
    ext = weather.get("q4_jours_extremes", {})
    rows = []
    for k, v in ext.items():
        if isinstance(v, dict) and v:
            m = [kk for kk in v if kk != "day"]
            rows.append([k.replace("_", " "), v.get("day"), v.get(m[0]) if m else ""])
    L.append(_md_table(["Événement", "Date", "Valeur"], rows))

    # Partie 2
    L.append("## Partie 2 — Flux Streaming : mobilité (vélos CityBikes)\n")
    _nsta = f"{mobility.get('n_stations', 0):,}".replace(",", " ")
    L.append(f"Périmètre : {_nsta} stations, {mobility.get('n_reseaux', 0)} réseaux.\n")
    L.append("### Stations les plus utilisées ?\n")
    L.append(_md_table(["Ville", "Station", "Occupation", "Capacité"],
             [[s.get("city"), s.get("station_name"), s.get("occupancy_rate"), s.get("capacity")]
              for s in mobility.get("q1_stations_plus_utilisees", [])[:6]]))
    L.append("### Zones à offre insuffisante ?\n")
    L.append(_md_table(["Ville", "Station", "% vide", "Vélos dispo"],
             [[s.get("city"), s.get("station_name"), s.get("pct_empty"), s.get("bikes_available")]
              for s in mobility.get("q2_offre_insuffisante", [])[:6]]))
    L.append("### Pics d'utilisation journaliers ?\n")
    L.append("Profil horaire de disponibilité (voir `mobility_profil_horaire.png`).\n")
    L.append("### Déséquilibres géographiques ?\n")
    L.append(_md_table(["Ville", "Stations", "Occupation", "% vides", "% pleines"],
             [[d.get("city"), d.get("stations"), d.get("occupancy_moy"),
               d.get("pct_empty_moy"), d.get("pct_full_moy")]
              for d in mobility.get("q4_desequilibres_villes", [])[:6]]))
    L.append("### Stations critiques à rééquilibrer ?\n")
    L.append(_md_table(["Ville", "Station", "Type", "Criticité"],
             [[s.get("city"), s.get("station_name"), s.get("type_critique"), s.get("criticite")]
              for s in mobility.get("q5_stations_critiques", [])[:6]]))

    # Partie 3
    L.append("## Partie 3 — Flux IoT : pollution (OpenAQ)\n")
    _nmes = f"{pollution.get('n_mesures', 0):,}".replace(",", " ")
    L.append(f"Périmètre : {_nmes} mesures, {pollution.get('n_villes', 0)} villes.\n")
    L.append("### Villes les plus polluées par polluant ?\n")
    L.append(_md_table(["Polluant", "Ville", "Valeur (µg/m³)"],
             [[k, v.get("ville"), round(v.get("valeur", 0), 1)]
              for k, v in pollution.get("q1_ville_plus_polluee_par_polluant", {}).items()]))
    L.append("### Dépassements des seuils OMS ?\n")
    dep = pollution.get("q3_depassements_seuils", {})
    L.append(f"Total : **{dep.get('total_depassements', 0)}** dépassements.\n")
    L.append(_md_table(["Ville", "Polluant", "Nb dépassements"],
             [[d.get("city"), d.get("parameter"), d.get("nb_depassements")]
              for d in dep.get("top", [])[:6]]))
    L.append("### Indice de qualité de l'air par ville ?\n")
    aqi = pollution.get("q4_indice_qualite_air", {})
    L.append(_md_table(["Ville", "Indice (ratio / seuil)"],
             [[c.get("city"), c.get("indice")] for c in aqi.get("classement", [])[:8]]))

    # Partie 4
    L.append("## Partie 4 — Analyse croisée\n")
    L.append(f"_{cross.get('methode', '')}_\n")
    L.append("### Relation météo ↔ pollution ?\n")
    q1c = cross.get("q1_meteo_pollution", {})
    if q1c.get("disponible"):
        corr = q1c.get("correlations", {})
        headers = ["Polluant"] + sorted({k for v in corr.values() for k in v})
        rows = [[pol] + [cc.get(h, "") for h in headers[1:]] for pol, cc in corr.items()]
        L.append(_md_table(headers, rows))
        o3 = corr.get("o3", {})
        if o3:
            L.append(f"**Résultat marquant : O₃ ↔ température = {o3.get('temperature_c')}** "
                     f"(formation photochimique de l'ozone).\n")
    L.append("### La météo influence-t-elle l'usage des vélos ?\n")
    q2c = cross.get("q2_meteo_velos", {})
    if q2c.get("disponible"):
        rows = [[t, k, v] for t, cc in q2c.get("correlations", {}).items() for k, v in cc.items()]
        L.append(_md_table(["Cible", "Variable météo", "Corrélation"], rows))
    else:
        L.append(f"{q2c.get('message', 'Nécessite une collecte vélos sur toute la journée.')}\n")
    L.append("### Conditions favorables à la mobilité douce ?\n")
    q3c = cross.get("q3_conditions_favorables", {})
    L.append(f"Définition : {q3c.get('definition_conditions_favorables', {})}. "
             f"Part d'heures favorables : **{q3c.get('part_heures_favorables_pct', 0)} %**.\n")
    saison = q3c.get("part_favorable_par_saison_pct", {})
    if saison:
        L.append(_md_table(["Saison", "Heures favorables (%)"], [[k, v] for k, v in saison.items()]))
    L.append("### Interactions météo × pollution × mobilité ?\n")
    q4c = cross.get("q4_triple_interaction", {})
    L.append(f"Heures communes analysées : {q4c.get('n_heures_communes', 0)} "
             f"(voir `cross_cycle_diurne.png`).\n")

    output = output or (config.PROJECT_ROOT / "docs" / "REPONSES_QUESTIONS_METIER.md")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(L), encoding="utf-8")
    log.info("Document Markdown généré : %s", output)
    return output


if __name__ == "__main__":
    build()
